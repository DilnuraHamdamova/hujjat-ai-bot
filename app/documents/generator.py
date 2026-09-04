from dataclasses import dataclass
from hashlib import sha256
from pathlib import Path
from uuid import UUID

from docx import Document as create_document
from docx.document import Document as DocxDocument
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Mm, Pt, RGBColor
from jinja2 import Environment, FileSystemLoader, select_autoescape

from app.services.localization import DOCUMENT_LABELS, OBJECTIVE_LABELS, normalize_language
from app.services.resume_schema import ResumeData
from app.services.templates import template_family


@dataclass(frozen=True)
class GeneratedArtifact:
    format: str
    path: Path
    checksum: str


class DocumentGenerator:
    def __init__(self, storage_dir: Path) -> None:
        self.storage_dir = storage_dir
        self.template_dir = Path(__file__).resolve().parents[1] / "templates"

    def generate(
        self,
        resume_id: UUID,
        raw_data: dict[str, object],
        language: str = "uz",
        template_code: str = "classic",
        file_format: str | None = None,
    ) -> list[GeneratedArtifact]:
        locale = normalize_language(language)
        output_dir = self.storage_dir / str(resume_id)
        output_dir.mkdir(parents=True, exist_ok=True)
        document_type = str(raw_data.get("document_type", "cv"))
        basename = "obyektivka" if document_type == "objective" else "cv"
        formats = (file_format,) if file_format in ("pdf", "docx") else ("docx", "pdf")
        artifacts: list[GeneratedArtifact] = []
        for selected_format in formats:
            path = output_dir / f"{basename}.{selected_format}"
            if document_type == "objective":
                if selected_format == "docx":
                    self._write_objective_docx(raw_data, path, locale)
                else:
                    self._write_objective_pdf(raw_data, path, locale)
            else:
                data = ResumeData.model_validate(raw_data)
                if selected_format == "docx":
                    self._write_resume_docx(data, raw_data, path, locale, template_code)
                else:
                    self._write_resume_pdf(data, raw_data, path, locale, template_code)
            artifacts.append(self._artifact(selected_format, path))
        return artifacts

    @staticmethod
    def _artifact(file_format: str, path: Path) -> GeneratedArtifact:
        return GeneratedArtifact(
            format=file_format,
            path=path,
            checksum=sha256(path.read_bytes()).hexdigest(),
        )

    @staticmethod
    def _normalize_objective_no_values(data: dict[str, object], language: str) -> dict[str, object]:
        normalized = dict(data)
        no_value = {"uz": "yo‘q", "en": "none", "ru": "нет"}[normalize_language(language)]
        for key in (
            "objective_party",
            "objective_degree",
            "objective_title",
            "objective_awards",
            "objective_elected",
        ):
            if str(normalized.get(key, "")).strip() == "-":
                normalized[key] = no_value
        return normalized

    @staticmethod
    def _write_resume_docx(
        data: ResumeData,
        raw_data: dict[str, object],
        path: Path,
        language: str,
        template_code: str,
    ) -> None:
        labels = DOCUMENT_LABELS[normalize_language(language)]
        document = create_document()
        styles = document.styles
        styles["Normal"].font.name = "Arial"
        styles["Normal"].font.size = Pt(10)

        # Keep the Word export visually distinct as well as the PDF export.
        section = document.sections[0]
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        family = template_family(template_code)
        if family == "modern":
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)
        elif family == "europass":
            section.left_margin = Inches(0.65)
            section.right_margin = Inches(0.65)

        photo_path = Path(str(raw_data.get("photo_path", "")))
        if photo_path.is_file():
            photo = document.add_paragraph()
            photo.alignment = (
                WD_ALIGN_PARAGRAPH.RIGHT
                if family in ("modern", "europass")
                else WD_ALIGN_PARAGRAPH.CENTER
            )
            photo.add_run().add_picture(str(photo_path), width=Inches(1.18), height=Inches(1.57))

        title_alignment = (
            WD_ALIGN_PARAGRAPH.LEFT
            if family in ("modern", "europass")
            else WD_ALIGN_PARAGRAPH.CENTER
        )
        title = document.add_paragraph()
        title.alignment = title_alignment
        name_run = title.add_run(data.full_name or "CV")
        name_run.bold = True
        name_run.font.size = Pt(22)
        if family == "europass":
            name_run.font.color.rgb = RGBColor(28, 83, 145)
        elif family == "modern":
            name_run.font.color.rgb = RGBColor(13, 148, 136)

        position = document.add_paragraph()
        position.alignment = title_alignment
        position_run = position.add_run(data.job_title or "")
        position_run.font.size = Pt(13)

        contacts = " | ".join(item for item in (data.phone, data.email, data.location) if item)
        contact_paragraph = document.add_paragraph(contacts)
        contact_paragraph.alignment = title_alignment

        DocumentGenerator._add_section(
            document, labels["profile"], [data.summary] if data.summary else []
        )
        DocumentGenerator._add_section(document, labels["skills"], data.skills)
        DocumentGenerator._add_section(document, labels["experience"], data.experience)
        DocumentGenerator._add_section(document, labels["education"], data.education)
        DocumentGenerator._add_section(document, labels["languages"], data.languages)
        DocumentGenerator._add_custom_docx_sections(document, raw_data)
        document.save(str(path))

    @staticmethod
    def _add_section(document: DocxDocument, title: str, values: list[str]) -> None:
        if not values:
            return
        heading = document.add_paragraph()
        heading_run = heading.add_run(title)
        heading_run.bold = True
        heading_run.font.size = Pt(12)
        for value in values:
            document.add_paragraph(value, style="List Bullet" if len(values) > 1 else None)

    @staticmethod
    def _add_custom_docx_sections(document: DocxDocument, raw_data: dict[str, object]) -> None:
        raw_sections = raw_data.get("custom_sections", [])
        if not isinstance(raw_sections, list):
            return
        for section in raw_sections:
            if not isinstance(section, dict):
                continue
            title = str(section.get("title", "")).strip()
            content = str(section.get("content", "")).strip()
            if title and content:
                DocumentGenerator._add_section(document, title.upper(), [content])

    def _environment(self) -> Environment:
        return Environment(
            loader=FileSystemLoader(self.template_dir),
            autoescape=select_autoescape(("html", "xml")),
        )

    def _write_resume_pdf(
        self,
        data: ResumeData,
        raw_data: dict[str, object],
        path: Path,
        language: str,
        template_code: str,
    ) -> None:
        from weasyprint import HTML  # type: ignore[import-untyped]

        labels = DOCUMENT_LABELS[normalize_language(language)]
        templates = {
            "classic": "resume.html",
            "classic_1": "resume.html",
            "classic_2": "resume.html",
            "classic_3": "resume.html",
            "modern": "resume_modern.html",
            "modern_1": "resume_modern.html",
            "modern_2": "resume_modern.html",
            "modern_3": "resume_modern.html",
            "europass": "resume_europass.html",
            "europass_1": "resume_europass.html",
            "europass_2": "resume_europass.html",
            "europass_3": "resume_europass.html",
        }
        template_name = templates.get(template_code, "resume.html")
        photo_path = Path(str(raw_data.get("photo_path", "")))
        html = (
            self._environment()
            .get_template(template_name)
            .render(
                resume=data,
                language=language,
                labels=labels,
                template_code=template_code,
                custom_sections=raw_data.get("custom_sections", []),
                photo_uri=photo_path.resolve().as_uri() if photo_path.is_file() else "",
            )
        )
        HTML(string=html, base_url=str(self.template_dir)).write_pdf(path)

    def _write_objective_docx(self, data: dict[str, object], path: Path, language: str) -> None:
        data = self._normalize_objective_no_values(data, language)
        labels = OBJECTIVE_LABELS[normalize_language(language)]
        document = create_document()
        section = document.sections[0]
        section.top_margin = Inches(0.45)
        section.bottom_margin = Inches(0.45)
        section.left_margin = Inches(0.55)
        section.right_margin = Inches(0.55)
        document.styles["Normal"].font.name = "Times New Roman"
        document.styles["Normal"].font.size = Pt(11)
        document.styles["Normal"].paragraph_format.space_after = Pt(0)

        header = document.add_table(rows=1, cols=2)
        header.autofit = False
        header.columns[0].width = Inches(5.8)
        header.columns[1].width = Inches(1.25)
        identity_cell, photo_cell = header.rows[0].cells

        title = identity_cell.paragraphs[0]
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title.add_run(labels["document_title"])
        title_run.bold = True
        title_run.font.size = Pt(15)

        full_name = str(data.get("objective_full_name", ""))
        heading = identity_cell.add_paragraph()
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        heading_run = heading.add_run(full_name.upper())
        heading_run.bold = True
        heading_run.font.size = Pt(12)

        photo_path = Path(str(data.get("photo_path", "")))
        if photo_path.is_file():
            photo_paragraph = photo_cell.paragraphs[0]
            photo_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            photo_paragraph.add_run().add_picture(
                str(photo_path), width=Inches(1.18), height=Inches(1.57)
            )

        paired_fields = (
            (("objective_birth_date", "birth_date"), ("objective_birth_place", "birth_place")),
            (("objective_nationality", "nationality"), ("objective_party", "party")),
            (
                ("objective_education_level", "education_level"),
                ("objective_graduated", "graduated"),
            ),
        )
        details = document.add_table(rows=0, cols=4)
        for left, right in paired_fields:
            row = details.add_row().cells
            for offset, (key, label_key) in zip((0, 2), (left, right), strict=True):
                label_run = row[offset].paragraphs[0].add_run(f"{labels[label_key]}:")
                label_run.bold = True
                value = data.get(key, "")
                rendered = "\n".join(map(str, value)) if isinstance(value, list) else str(value)
                row[offset + 1].text = rendered

        specialty_table = document.add_table(rows=0, cols=2)
        if "objective_specialty" in data:
            cells = specialty_table.add_row().cells
            specialty_run = cells[0].paragraphs[0].add_run(f"{labels['specialty']}:")
            specialty_run.bold = True
            cells[1].text = str(data["objective_specialty"])

        degree_table = document.add_table(rows=1, cols=4)
        for offset, (key, label_key) in zip(
            (0, 2),
            (("objective_degree", "degree"), ("objective_title", "academic_title")),
            strict=True,
        ):
            cells = degree_table.rows[0].cells
            label_run = cells[offset].paragraphs[0].add_run(f"{labels[label_key]}:")
            label_run.bold = True
            cells[offset + 1].text = str(data.get(key, ""))

        full_width_fields = (
            ("objective_languages", "foreign_languages"),
            ("objective_awards", "awards"),
            ("objective_elected", "elected"),
        )
        full_details = document.add_table(rows=0, cols=2)
        for key, label_key in full_width_fields:
            if key not in data:
                continue
            cells = full_details.add_row().cells
            label_run = cells[0].paragraphs[0].add_run(f"{labels[label_key]}:")
            label_run.bold = True
            value = data[key]
            cells[1].text = "\n".join(map(str, value)) if isinstance(value, list) else str(value)

        employment = data.get("objective_employment")
        if isinstance(employment, list) and employment:
            employment_heading = document.add_paragraph()
            employment_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            employment_run = employment_heading.add_run(labels["employment"])
            employment_run.bold = True
            employment_table = document.add_table(rows=0, cols=2)
            employment_table.autofit = False
            employment_table.columns[0].width = Inches(1.55)
            employment_table.columns[1].width = Inches(5.45)
            for item in map(str, employment):
                parts = [part.strip() for part in item.split("|")]
                cells = employment_table.add_row().cells
                cells[0].width = Inches(1.55)
                cells[1].width = Inches(5.45)
                if len(parts) >= 2:
                    cells[0].text = parts[0]
                    cells[1].text = ", ".join(part for part in parts[1:] if part)
                    for cell in cells:
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                run.font.size = Pt(12)
                else:
                    cells[0].merge(cells[1]).text = item
                    for run in cells[0].paragraphs[0].runs:
                        run.font.size = Pt(12)
        DocumentGenerator._add_custom_docx_sections(document, data)

        relatives = data.get("objective_relatives")
        if isinstance(relatives, list) and relatives:
            relatives_section = document.add_section(WD_SECTION.NEW_PAGE)
            relatives_section.page_width = Mm(210)
            relatives_section.page_height = Mm(297)
            relatives_section.top_margin = Inches(0.45)
            relatives_section.bottom_margin = Inches(0.45)
            relatives_section.left_margin = Inches(0.55)
            relatives_section.right_margin = Inches(0.55)
            DocumentGenerator._add_relatives_docx(
                document, list(map(str, relatives)), labels, full_name
            )
        document.save(str(path))

    @staticmethod
    def _add_relatives_docx(
        document: DocxDocument,
        relatives: list[str],
        labels: dict[str, str],
        full_name: str,
    ) -> None:
        heading = document.add_paragraph()
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = heading.add_run(f"{full_name}ning yaqin qarindoshlari to‘g‘risida\nMA’LUMOT")
        run.bold = True
        table = document.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        headers = (
            labels["relationship"],
            labels["relative_name"],
            labels["relative_birth"],
            labels["relative_work"],
            labels["relative_address"],
        )
        for cell, header in zip(table.rows[0].cells, headers, strict=True):
            cell.text = header
        for relative in relatives:
            parts = [part.strip() for part in relative.split("|")]
            parts.extend([""] * (5 - len(parts)))
            cells = table.add_row().cells
            for cell, value in zip(cells, parts[:5], strict=True):
                cell.text = value

    def _write_objective_pdf(self, data: dict[str, object], path: Path, language: str) -> None:
        from weasyprint import HTML

        data = self._normalize_objective_no_values(data, language)
        labels = OBJECTIVE_LABELS[normalize_language(language)]
        photo_path = Path(str(data.get("photo_path", "")))
        relatives = []
        raw_relatives = data.get("objective_relatives", [])
        if isinstance(raw_relatives, list):
            for relative in raw_relatives:
                parts = [part.strip() for part in str(relative).split("|")]
                parts.extend([""] * (5 - len(parts)))
                relatives.append(parts[:5])
        html = (
            self._environment()
            .get_template("objective.html")
            .render(
                data=data,
                language=language,
                labels=labels,
                photo_uri=photo_path.resolve().as_uri() if photo_path.is_file() else "",
                relatives=relatives,
                custom_sections=data.get("custom_sections", []),
            )
        )
        HTML(string=html, base_url=str(self.template_dir)).write_pdf(path)
